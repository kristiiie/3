"""Создание DOCX-шаблона с полями слияния Word.

Шаблон содержит поля MERGEFIELD, которые можно связать с CSV-файлом через
раздел «Рассылки» в Microsoft Word.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_TEMPLATE = BASE_DIR / "award_template_merge_fields.docx"


def configure_document(document: Document) -> None:
    """Настраивает страницу и рамку шаблона."""
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(14)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    section_properties = section._sectPr
    page_borders = OxmlElement("w:pgBorders")
    page_borders.set(qn("w:offsetFrom"), "page")
    for border_name in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "18")
        border.set(qn("w:space"), "18")
        border.set(qn("w:color"), "1F4E78")
        page_borders.append(border)
    section_properties.append(page_borders)


def add_merge_field(paragraph, field_name: str, size=14, bold=False, italic=False) -> None:
    """Вставляет поле MERGEFIELD в текущий абзац с заданным оформлением."""
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f"MERGEFIELD {field_name} \\* MERGEFORMAT")

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_fonts = OxmlElement("w:rFonts")
    run_fonts.set(qn("w:ascii"), "Times New Roman")
    run_fonts.set(qn("w:hAnsi"), "Times New Roman")
    run_fonts.set(qn("w:eastAsia"), "Times New Roman")
    run_properties.append(run_fonts)

    run_size = OxmlElement("w:sz")
    run_size.set(qn("w:val"), str(size * 2))
    run_properties.append(run_size)
    if bold:
        run_properties.append(OxmlElement("w:b"))
    if italic:
        run_properties.append(OxmlElement("w:i"))
    run.append(run_properties)

    text = OxmlElement("w:t")
    text.text = f"«{field_name}»"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def add_text(paragraph, text: str, size=14, bold=False, italic=False) -> None:
    """Добавляет текстовый фрагмент с базовым шрифтом."""
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_template_paragraph(document: Document, parts, size=14, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=5):
    """Добавляет абзац шаблона. parts: строки и кортежи ('field', имя_поля)."""
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    for part in parts:
        if isinstance(part, tuple) and part[0] == "field":
            add_merge_field(paragraph, part[1], size=size, bold=bold, italic=italic)
        else:
            add_text(paragraph, str(part), size=size, bold=bold, italic=italic)
    return paragraph


def add_separator(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("―" * 42)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def build_template() -> None:
    """Формирует файл шаблона с полями слияния."""
    document = Document()
    configure_document(document)
    add_template_paragraph(document, [("field", "organizer")], size=13, bold=True, space_after=6)
    add_separator(document)
    add_template_paragraph(document, [("field", "title_text")], size=30, bold=True, space_after=8)
    add_template_paragraph(document, [("field", "contest_name")], size=17, italic=True, space_after=14)
    add_template_paragraph(document, [("field", "award_text")], size=15, space_after=4)
    add_template_paragraph(document, [("field", "recipient_name")], size=23, bold=True, space_after=4)
    add_template_paragraph(document, [("field", "recipient_role")], size=15, italic=True, space_after=12)
    add_template_paragraph(document, ["Номинация: «", ("field", "nomination"), "». Возрастная группа: ", ("field", "age_group"), "."], size=15, space_after=8)
    add_template_paragraph(document, ["Конкурсная работа: «", ("field", "work_title"), "»"], size=15, space_after=8)
    add_template_paragraph(document, [("field", "note")], size=13, italic=True, space_after=16)
    add_separator(document)
    add_template_paragraph(document, [("field", "signature_role"), " ____________________ ", ("field", "signatory_name")], size=13, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=3)
    add_template_paragraph(document, [("field", "city"), ", ", ("field", "issue_date")], size=13, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    add_template_paragraph(document, ["Регистрационный номер: ", ("field", "document_id")], size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    document.save(OUTPUT_TEMPLATE)


if __name__ == "__main__":
    build_template()
    print(f"Создан шаблон: {OUTPUT_TEMPLATE}")
