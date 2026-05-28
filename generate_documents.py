"""Генерация наградных документов из CSV-таблицы.

Скрипт читает файл awards_data.csv с разделителем ';' и формирует один DOCX-файл,
в котором каждая строка таблицы превращается в отдельный наградной документ.
PDF можно получить в Microsoft Word через команду «Сохранить как PDF» или
автоматически через LibreOffice, если он установлен в системе.
"""

from __future__ import annotations

import csv
import shutil
import subprocess
from pathlib import Path
from typing import Dict, Iterable, List

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE_DIR = Path(__file__).resolve().parent
DATA_FILE = BASE_DIR / "awards_data.csv"
OUTPUT_DOCX = BASE_DIR / "award_documents_merged.docx"


def load_rows(csv_path: Path) -> List[Dict[str, str]]:
    """Загружает строки CSV-файла с данными наградных документов."""
    with csv_path.open("r", encoding="utf-8-sig", newline="") as file_obj:
        reader = csv.DictReader(file_obj, delimiter=";")
        return [dict(row) for row in reader]


def configure_document(document: Document) -> None:
    """Настраивает страницу, поля, шрифт и рамку документа."""
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(14)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    section_properties = section._sectPr
    page_borders = OxmlElement("w:pgBorders")
    page_borders.set(qn("w:offsetFrom"), "page")
    for border_name in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "18")
        border.set(qn("w:space"), "18")
        border.set(qn("w:color"), "1F4E78")
        page_borders.append(border)
    section_properties.append(page_borders)


def add_paragraph(
    document: Document,
    text: str,
    size: int = 14,
    bold: bool = False,
    italic: bool = False,
    align: int = WD_ALIGN_PARAGRAPH.CENTER,
    space_after: int = 4,
) -> None:
    """Добавляет абзац с единым оформлением."""
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_separator(document: Document) -> None:
    """Добавляет декоративную горизонтальную линию."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("―" * 42)
    run.font.size = Pt(11)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    paragraph.paragraph_format.space_after = Pt(4)


def add_award_page(document: Document, row: Dict[str, str], is_last: bool = False) -> None:
    """Создает одну страницу наградного документа по строке CSV."""
    add_paragraph(document, row["organizer"], size=13, bold=True, space_after=6)
    add_separator(document)
    add_paragraph(document, row["title_text"], size=30, bold=True, space_after=8)
    add_paragraph(document, row["contest_name"], size=17, italic=True, space_after=14)
    add_paragraph(document, row["award_text"], size=15, space_after=4)
    add_paragraph(document, row["recipient_name"], size=23, bold=True, space_after=4)
    add_paragraph(document, row["recipient_role"], size=15, italic=True, space_after=12)

    if row["document_type"] == "Благодарственное письмо":
        body = (
            f"за профессиональное сопровождение творческой работы обучающегося, "
            f"поддержку детской инициативы и подготовку участника конкурса: {row['child_name']}."
        )
    elif row["document_type"] == "Диплом":
        body = (
            f"за {row['place']} в номинации «{row['nomination']}» "
            f"в возрастной группе {row['age_group']}."
        )
    else:
        body = (
            f"за участие в номинации «{row['nomination']}» "
            f"в возрастной группе {row['age_group']}."
        )
    add_paragraph(document, body, size=15, space_after=8)
    add_paragraph(document, f"Конкурсная работа: «{row['work_title']}»", size=15, space_after=8)
    add_paragraph(document, row["note"], size=13, italic=True, space_after=16)

    add_separator(document)
    signature = f"{row['signature_role']} ____________________ {row['signatory_name']}"
    add_paragraph(document, signature, size=13, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=3)
    add_paragraph(document, f"{row['city']}, {row['issue_date']}", size=13, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    add_paragraph(document, f"Регистрационный номер: {row['document_id']}", size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)

    if not is_last:
        document.add_page_break()


def build_awards(rows: Iterable[Dict[str, str]], output_path: Path) -> None:
    """Формирует итоговый DOCX-файл с наградными документами."""
    row_list = list(rows)
    document = Document()
    configure_document(document)
    for index, row in enumerate(row_list):
        add_award_page(document, row, is_last=index == len(row_list) - 1)
    document.save(output_path)


def convert_to_pdf(docx_path: Path) -> Path | None:
    """Пробует преобразовать DOCX в PDF через LibreOffice."""
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not libreoffice:
        return None
    subprocess.run(
        [
            libreoffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(docx_path.parent),
            str(docx_path),
        ],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    return docx_path.with_suffix(".pdf")


def main() -> None:
    """Точка входа: чтение CSV, формирование DOCX и попытка экспорта в PDF."""
    rows = load_rows(DATA_FILE)
    build_awards(rows, OUTPUT_DOCX)
    pdf_path = convert_to_pdf(OUTPUT_DOCX)
    print(f"Сформирован DOCX: {OUTPUT_DOCX}")
    if pdf_path:
        print(f"Сформирован PDF: {pdf_path}")
    else:
        print("PDF не создан автоматически: LibreOffice не найден. Сохраните файл как PDF из Word.")


if __name__ == "__main__":
    main()
